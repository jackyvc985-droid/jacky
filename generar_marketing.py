#!/usr/bin/env python3
"""Material de ventas: presentación de 10 láminas (HTML + PDF), infografías 1080x1350, mensajes de WhatsApp y guion.
Personalización: assets/logo.png y assets/config_marca.json."""
import base64
import json
import os
import shutil
import subprocess

RAIZ = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(RAIZ, "marketing")
CH = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
CFG = json.load(open(os.path.join(RAIZ, "assets", "config_marca.json"), encoding="utf-8"))
ETIQUETA = CFG.get("marca") or "SISTEMA ADMINISTRATIVO"
LOGO = os.path.join(RAIZ, "assets", "logo.png")
PREV = os.path.join(RAIZ, "vistas_previas")

NAVY, TEAL, CORAL, GOLD, GRIS = "#0B2A4A", "#0E7C8B", "#E8604C", "#C9A227", "#EEF2F5"

CSS = """
*{box-sizing:border-box;margin:0;padding:0}
:root{--navy:#0B2A4A;--teal:#0E7C8B;--coral:#E8604C;--gold:#C9A227;--gris:#EEF2F5;--tx:#1c2b3a}
body{font-family:Abadi,'Abadi MT',Inter,'Segoe UI',Arial,sans-serif;color:var(--tx);background:#071d35}
.dark{background:linear-gradient(135deg,#071d35 0%,#0B2A4A 55%,#134a6c 100%);color:#fff}
.light{background:#fff;color:var(--tx)}
.wm{font-weight:800;letter-spacing:.14em;font-size:15px;color:var(--gold);display:flex;align-items:center;gap:10px}
.wm:before{content:"";width:8px;height:26px;background:var(--gold);display:inline-block}
.light .wm{color:var(--navy)}
.logo{height:34px;width:auto;object-fit:contain}
.kicker{font-size:14px;font-weight:800;letter-spacing:.16em;color:var(--gold);text-transform:uppercase}
.light .kicker{color:var(--teal)}
h1{font-size:68px;line-height:1.04;font-weight:800;letter-spacing:-.01em}
h2{font-size:40px;line-height:1.12;font-weight:800;letter-spacing:-.01em;color:var(--navy)}
.dark h2{color:#fff}
.sub{font-size:20px;line-height:1.4;opacity:.85}
.chip{display:inline-block;border:2px solid var(--gold);color:var(--gold);border-radius:999px;padding:8px 18px;font-size:13px;font-weight:800;letter-spacing:.08em;margin-right:10px}
.card{background:var(--gris);border-radius:16px;padding:22px 24px}
.dark .card{background:rgba(255,255,255,.08);border:1px solid rgba(255,255,255,.14)}
.card h3{font-size:19px;color:var(--navy);margin:10px 0 6px}
.dark .card h3{color:#fff}
.card p,.card li{font-size:15px;line-height:1.4}
ul{padding-left:18px}
.ico{width:46px;height:46px;border-radius:12px;display:flex;align-items:center;justify-content:center;background:var(--navy);color:#fff}
.ico svg{width:26px;height:26px;stroke:currentColor;fill:none;stroke-width:2;stroke-linecap:round;stroke-linejoin:round}
.dark .ico{background:var(--gold);color:var(--navy)}
.flow{display:flex;align-items:stretch;gap:0}
.node{flex:1;background:var(--gris);padding:12px 12px 12px 22px;font-size:14px;font-weight:700;color:var(--navy);position:relative;clip-path:polygon(0 0,calc(100% - 16px) 0,100% 50%,calc(100% - 16px) 100%,0 100%,16px 50%);margin-left:-8px;display:flex;align-items:center;justify-content:center;text-align:center;min-height:54px}
.node:first-child{clip-path:polygon(0 0,calc(100% - 16px) 0,100% 50%,calc(100% - 16px) 100%,0 100%);margin-left:0}
.node.a{background:var(--navy);color:#fff}.node.b{background:var(--teal);color:#fff}.node.c{background:var(--gold);color:var(--navy)}.node.d{background:var(--coral);color:#fff}
table{border-collapse:collapse;width:100%}
th{background:var(--navy);color:#fff;text-align:left;padding:12px 14px;font-size:13px;letter-spacing:.06em;text-transform:uppercase}
td{padding:11px 14px;font-size:14px;border-bottom:1px solid #dbe3ea}
tr:nth-child(even) td{background:#f6f8fa}
td:first-child{font-weight:800;color:var(--navy)}
.sem{display:inline-block;width:12px;height:12px;border-radius:50%;margin-right:8px}
.v{background:#3DAE6B}.am{background:var(--gold)}.r{background:var(--coral)}
"""

LOGO_B64 = None
if os.path.exists(LOGO):
    LOGO_B64 = "data:image/png;base64," + base64.b64encode(open(LOGO, "rb").read()).decode()


def marca():
    if LOGO_B64:
        return f'<img class="logo" src="{LOGO_B64}" alt="{ETIQUETA}">'
    return f'<div class="wm">{ETIQUETA}</div>'


def contacto(sep=" · "):
    datos = [CFG.get(k, "") for k in ("correo", "telefono", "web")]
    return sep.join(d for d in datos if d)


def img64(ruta):
    ruta = os.path.join(PREV, ruta)
    if not os.path.exists(ruta):
        return ""
    return "data:image/png;base64," + base64.b64encode(open(ruta, "rb").read()).decode()


ICONOS = {
    "doc": '<svg viewBox="0 0 24 24"><path d="M6 3h8l4 4v14H6z"/><path d="M14 3v4h4M9 12h6M9 16h6"/></svg>',
    "loop": '<svg viewBox="0 0 24 24"><path d="M20 12a8 8 0 1 1-2.3-5.6"/><path d="M20 4v5h-5"/></svg>',
    "shield": '<svg viewBox="0 0 24 24"><path d="M12 3l8 3v6c0 5-3.5 8-8 9-4.5-1-8-4-8-9V6z"/><path d="M9 12l2 2 4-4"/></svg>',
    "chart": '<svg viewBox="0 0 24 24"><path d="M4 20V4M4 20h16"/><path d="M8 16v-4M12 16V8M16 16v-6"/></svg>',
    "people": '<svg viewBox="0 0 24 24"><circle cx="9" cy="8" r="3"/><path d="M3 20c0-3.5 2.7-6 6-6s6 2.5 6 6"/><path d="M16 5a3 3 0 0 1 0 6M21 20c0-2.6-1.5-4.6-3.5-5.5"/></svg>',
    "cart": '<svg viewBox="0 0 24 24"><path d="M3 4h2l2.5 11h10L20 7H6"/><circle cx="9" cy="19" r="1.5"/><circle cx="17" cy="19" r="1.5"/></svg>',
    "cash": '<svg viewBox="0 0 24 24"><rect x="3" y="6" width="18" height="12" rx="2"/><circle cx="12" cy="12" r="2.5"/></svg>',
    "check": '<svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="9"/><path d="M8 12l3 3 5-6"/></svg>',
    "flag": '<svg viewBox="0 0 24 24"><path d="M5 21V4M5 4h11l-2 4 2 4H5"/></svg>',
    "cal": '<svg viewBox="0 0 24 24"><rect x="4" y="5" width="16" height="15" rx="2"/><path d="M4 10h16M9 3v4M15 3v4"/></svg>',
}


def ico(n):
    return f'<div class="ico">{ICONOS[n]}</div>'


# ----------------------------------------------------------------------------- PRESENTACIÓN
def pie(n, oscuro=False):
    return f'<div class="foot" style="position:absolute;left:64px;right:64px;bottom:26px;display:flex;justify-content:space-between;font-size:12px;opacity:.65"><span>{CFG["lema"]}</span><span>{n:02d} / 10</span></div>'


def top(oscuro=False, kicker=""):
    return (f'<div style="position:absolute;left:64px;right:64px;top:34px;display:flex;justify-content:space-between;align-items:center">{marca()}'
            f'<span class="kicker">{kicker}</span></div>')


def lamina(n, clase, cuerpo, kicker=""):
    return f'<section class="slide {clase}" data-n="{n}">{top(clase=="dark", kicker)}{cuerpo}{pie(n)}</section>'


def deck_html():
    s = []
    # 1 portada
    tablero = img64("7_Control_Directivo/01_Control_Operativo_Semanal_resumen.png")
    s.append(f'''<section class="slide dark" data-n="1">
      <div style="position:absolute;right:-140px;top:-160px;width:620px;height:620px;border-radius:50%;background:rgba(201,162,39,.14);border:2px solid rgba(201,162,39,.4)"></div>
      <div style="position:absolute;left:-120px;bottom:-200px;width:460px;height:460px;border-radius:50%;background:rgba(255,255,255,.05)"></div>
      {top(True, "Propuesta comercial")}
      <div style="position:absolute;left:64px;top:170px;width:700px">
        <div class="kicker" style="margin-bottom:18px">Sistema administrativo profesional</div>
        <h1>ORDEN PARA OPERAR.<br><span style="color:var(--gold)">CONTROL PARA CRECER.</span></h1>
        <p class="sub" style="margin-top:26px;max-width:600px">Formatos, controles y método para que constructoras, empresas industriales y de servicios operen con trazabilidad y cierren bien cada proyecto.</p>
        <div style="margin-top:34px"><span class="chip">NÓMINA</span><span class="chip">FISCAL E IMSS</span><span class="chip">COBRANZA</span><span class="chip">PROYECTOS</span></div>
      </div>
      <img src="{tablero}" style="position:absolute;right:50px;top:200px;width:430px;border-radius:10px;box-shadow:0 30px 60px rgba(0,0,0,.45);transform:rotate(4deg)">
      {pie(1)}</section>''')
    # 2 problema
    def pc(i, t, d):
        return f'<div class="card">{ico(i)}<h3>{t}</h3><p>{d}</p></div>'
    s.append(lamina(2, "light", f'''<div style="position:absolute;left:64px;right:64px;top:110px">
      <div class="kicker">El problema</div><h2 style="margin:10px 0 30px;max-width:900px">Cuando la administración improvisa, la operación paga la factura</h2>
      <div style="display:grid;grid-template-columns:repeat(4,1fr);gap:18px">
        {pc("doc","Documentos dispersos","Contratos, facturas y evidencias en correos, carpetas y WhatsApp; nadie sabe cuál es la versión final.")}
        {pc("loop","Retrabajos","Se repite lo que ya se hizo porque no hay procesos ni responsables definidos.")}
        {pc("shield","Riesgo fiscal y laboral","Nómina, impuestos e IMSS se atienden tarde o sin evidencia suficiente.")}
        {pc("chart","Decisiones a ciegas","Dirección no ve avances, pendientes ni riesgos hasta que ya son un problema.")}
      </div>
      <div style="margin-top:28px;background:var(--navy);color:#fff;border-radius:14px;padding:20px 26px;font-size:20px"><b style="color:var(--gold)">Resultado:</b> sobrecostos, multas, cobranza lenta y proyectos que se entregan pero no se cierran.</div>
    </div>''', "Problema"))
    # 3 solución
    s.append(lamina(3, "dark", f'''<div style="position:absolute;left:64px;right:64px;top:110px">
      <div class="kicker">La solución</div><h2 style="margin:10px 0 28px;max-width:940px">Un sistema listo para operar: formatos, controles y método</h2>
      <div style="display:grid;grid-template-columns:repeat(3,1fr);gap:20px">
        <div class="card">{ico("doc")}<h3>38 formatos profesionales</h3><p>Excel con fórmulas y tableros, y documentos Word para contratos, políticas, actas y reportes.</p></div>
        <div class="card">{ico("chart")}<h3>Controles con semáforos</h3><p>Seguimiento semanal de actividades, pendientes, riesgos, decisiones e indicadores.</p></div>
        <div class="card">{ico("flag")}<h3>Método de cuatro pasos</h3><p>Ordenar, planear, ejecutar y cerrar, con responsables, fechas y evidencia en cada paso.</p></div>
      </div>
      <div style="margin-top:30px;display:flex;gap:44px;font-size:18px">
        <div><b style="color:var(--gold);font-size:34px">↓</b><br>Menos retrabajo</div><div><b style="color:var(--gold);font-size:34px">↑</b><br>Control documental</div>
        <div><b style="color:var(--gold);font-size:34px">✓</b><br>Decisiones con información</div><div><b style="color:var(--gold);font-size:34px">⚑</b><br>Proyectos cerrados con evidencia</div></div>
    </div>''', "Propuesta de valor"))
    # 4 modelo
    paso = lambda cl, n, t, items, ent: (f'<div style="flex:1"><div class="node {cl}" style="min-height:64px;font-size:20px;margin-left:0">{n}. {t}</div>'
                                         f'<div class="card" style="margin:14px 8px 0 0;min-height:230px"><ul>{"".join(f"<li>{i}</li>" for i in items)}</ul>'
                                         f'<p style="margin-top:12px;font-size:12px;color:var(--teal);font-weight:800;letter-spacing:.08em">ENTREGABLES</p><p>{ent}</p></div></div>')
    s.append(lamina(4, "light", f'''<div style="position:absolute;left:64px;right:64px;top:110px">
      <div class="kicker">Método</div><h2 style="margin:10px 0 28px;max-width:900px">Cuatro pasos para pasar del desorden al control</h2>
      <div style="display:flex;gap:8px">
        {paso("a", 1, "ORDENAR", ["Levantar procesos y documentos", "Definir responsables", "Eliminar duplicidades"], "Mapa de procesos y expediente")}
        {paso("b", 2, "PLANEAR", ["Priorizar por urgencia e impacto", "Calendario y presupuesto", "Matriz de riesgos"], "Plan de trabajo y presupuesto")}
        {paso("c", 3, "EJECUTAR", ["Operar con formatos estándar", "Seguimiento semanal con semáforos", "Escalar desviaciones"], "Control semanal y reporte")}
        {paso("d", 4, "CERRAR", ["Validar entregables", "Conciliar cifras y pendientes", "Documentar lecciones"], "Acta de cierre y expediente")}
      </div>
      <p style="margin-top:22px;font-size:18px;color:var(--navy)"><b>Un proyecto no termina al entregar. Termina al cerrar bien.</b></p>
    </div>''', "Modelo operativo"))
    # 5 procesos
    def fl(tit, nodos, ev, cl):
        return (f'<div style="margin-bottom:14px"><div style="display:flex;align-items:center;gap:12px;margin-bottom:6px"><b style="color:var(--navy);width:210px;font-size:15px;letter-spacing:.05em">{tit}</b>'
                f'<div class="flow" style="flex:1">{"".join(f"<div class=node style=background:{cl[min(i, len(cl) - 1)][0]};color:{cl[min(i, len(cl) - 1)][1]}>{n}</div>" for i, n in enumerate(nodos))}</div></div>'
                f'<div style="margin-left:222px;font-size:12px;color:#5b6b78">Evidencia: {ev}</div></div>')
    cols = [("#0B2A4A", "#fff"), ("#0B2A4A", "#fff"), ("#0E7C8B", "#fff"), ("#0E7C8B", "#fff"), ("#C9A227", "#0B2A4A")]
    s.append(lamina(5, "light", f'''<div style="position:absolute;left:64px;right:64px;top:104px">
      <div class="kicker">Procesos clave</div><h2 style="margin:10px 0 24px;max-width:960px">Procesos visibles, con responsable, fecha y evidencia</h2>
      {fl("NÓMINA", ["Asistencia", "Incidencias", "Cálculo", "Dispersión", "Recibos y archivo"], "layout de dispersión, recibos CFDI, control de asistencia", cols)}
      {fl("FISCAL E IMSS", ["Documentos", "Cálculo", "Declaración", "Pago y acuse", "Expediente"], "acuses SAT e IDSE, línea de captura pagada", cols)}
      {fl("FACTURACIÓN Y COBRANZA", ["Solicitud", "Factura", "Seguimiento", "Cobro", "Conciliación"], "CFDI, estado de cuenta, conciliación bancaria", cols)}
      {fl("COMPRAS Y PROYECTOS", ["Requisición", "Cotización", "Autorización", "Orden y recepción", "Cierre"], "cuadro comparativo, orden de compra, acta de entrega-recepción", cols)}
    </div>''', "Diagramas de proceso"))
    # 6 tablero
    s.append(lamina(6, "dark", f'''<img src="{tablero}" style="position:absolute;left:64px;top:110px;width:520px;border-radius:10px;box-shadow:0 25px 55px rgba(0,0,0,.5)">
      <div style="position:absolute;left:640px;right:64px;top:112px"><div class="kicker">Tablero semanal</div>
      <h2 style="margin:10px 0 22px">Un tablero que dirección lee en un minuto</h2>
      <div class="card" style="margin-bottom:12px"><span class="sem v"></span><b>Verde:</b> en control. <span class="sem am" style="margin-left:12px"></span><b>Ámbar:</b> atención. <span class="sem r" style="margin-left:12px"></span><b>Coral:</b> acción inmediata.</div>
      <ul style="font-size:17px;line-height:1.7"><li>Actividades cumplidas, vencidas y pendientes abiertos</li><li>Riesgos con matriz de calor (probabilidad x impacto)</li><li>Decisiones que requiere dirección y fecha límite</li><li>Avance planeado vs. real por proyecto</li><li>Texto automático listo para el reporte semanal</li></ul></div>''', "Gráficas y tableros"))
    # 7 incluye
    col = lambda ico_, n, t, d: f'<div class="card" style="display:flex;gap:14px;align-items:flex-start">{ico(ico_)}<div><h3 style="margin:0 0 4px">{n} · {t}</h3><p>{d}</p></div></div>'
    s.append(lamina(7, "light", f'''<style>.s7 .card{{padding:13px 18px}}.s7 .card h3{{font-size:17px}}.s7 .card p{{font-size:13px}}.s7 .ico{{width:40px;height:40px}}</style><div class="s7" style="position:absolute;left:64px;right:64px;top:100px">
      <div class="kicker">Qué incluye</div><h2 style="margin:8px 0 16px;max-width:900px">38 formatos para ordenar toda la operación administrativa</h2>
      <div style="display:grid;grid-template-columns:repeat(2,1fr);gap:9px">
        {col("cash", 8, "Contabilidad y finanzas", "Ingresos y gastos, flujo, cobranza, conciliación, presupuesto, activos")}
        {col("people", 5, "Recursos humanos", "Asistencia, nómina, vacaciones, evaluación, permisos")}
        {col("cart", 4, "Inventarios y compras", "Inventario, kardex, orden de compra, cuadro comparativo")}
        {col("doc", 3, "Documentos y actas", "Cotización, recibo, minuta")}
        {col("shield", 8, "Contratos, cartas y actas", "Servicios, renuncia, constancia, cobranza, NDA, poder")}
        {col("check", 6, "Políticas y procedimientos", "Caja chica, compras, viáticos, asistencia, inventarios, alta")}
        {col("chart", 4, "Control directivo", "Control operativo semanal, reporte, plan de 30 días, modelo operativo")}
        <div class="card" style="background:var(--navy);color:#fff;display:flex;align-items:center"><div><b style="font-size:30px;color:var(--gold)">Excel + Word</b><br>Editables, con tu logotipo y listos para imprimir o enviar en PDF.</div></div>
      </div></div>''', "Entregables"))
    # 8 tabla
    sem = lambda c: f'<span class="sem {c}"></span>'
    s.append(lamina(8, "light", f'''<div style="position:absolute;left:64px;right:64px;top:106px">
      <div class="kicker">Control ejecutivo</div><h2 style="margin:10px 0 20px;max-width:960px">Qué se controla, quién responde y qué evidencia queda</h2>
      <table><tr><th>Área</th><th>Gestión</th><th>Control</th><th>Responsable</th><th>Evidencia</th></tr>
      <tr><td>Nómina</td><td>Cálculo y dispersión</td><td>Asistencia, incidencias y neto</td><td>RH / Finanzas</td><td>Layout, recibos CFDI</td></tr>
      <tr><td>Fiscal e IMSS</td><td>Declaraciones, altas y bajas</td><td>Calendario y acuses</td><td>Finanzas / RH</td><td>Acuses SAT e IDSE</td></tr>
      <tr><td>Facturación y cobranza</td><td>Emisión y seguimiento</td><td>Antigüedad de saldos</td><td>Administración</td><td>Estado de cuenta, conciliación</td></tr>
      <tr><td>Compras</td><td>De requisición a recepción</td><td>Cuadro comparativo, 3 vías</td><td>Compras / Dirección</td><td>Orden, entrada, factura</td></tr>
      <tr><td>Proyectos</td><td>Avance y presupuesto</td><td>Plan vs. real, riesgos</td><td>Proyectos</td><td>Reporte semanal, acta de cierre</td></tr></table>
      <p style="margin-top:18px;font-size:15px">{sem("v")}En control &nbsp; {sem("am")}Atención &nbsp; {sem("r")}Acción inmediata</p></div>''', "Tablas ejecutivas"))
    # 9 plan 30 días
    sm = lambda cl, t, items, ent: (f'<div style="flex:1"><div class="node {cl}" style="margin-left:0;min-height:56px;font-size:18px">{t}</div><div class="card" style="margin:12px 8px 0 0;min-height:220px">'
                                    f'<ul>{"".join(f"<li>{i}</li>" for i in items)}</ul><p style="margin-top:10px;font-size:12px;color:var(--teal);font-weight:800;letter-spacing:.08em">ENTREGABLE</p><p>{ent}</p></div></div>')
    s.append(lamina(9, "light", f'''<div style="position:absolute;left:64px;right:64px;top:106px">
      <div class="kicker">Implementación</div><h2 style="margin:10px 0 24px;max-width:960px">En 30 días: de la solicitud al primer reporte para dirección</h2>
      <div style="display:flex;gap:8px">
        {sm("a", "SEMANA 1 · ORDENAR", ["Reunión de arranque", "Levantamiento de procesos", "Recepción de documentos"], "Informe de diagnóstico")}
        {sm("b", "SEMANA 2 · PLANEAR", ["Priorizar brechas", "Responsables y calendario", "Configurar libro de control"], "Plan firmado")}
        {sm("c", "SEMANA 3 · EJECUTAR", ["Operar con los controles", "Regularizar pendientes críticos", "Capacitar al equipo"], "Primer reporte semanal")}
        {sm("d", "SEMANA 4 · CERRAR", ["Revisar indicadores", "Documentar procedimientos", "Plan de mejora a 90 días"], "Expediente final")}
      </div></div>''', "Plan de 30 días"))
    # 10 cierre
    s.append(f'''<section class="slide dark" data-n="10">
      <div style="position:absolute;right:-160px;bottom:-200px;width:640px;height:640px;border-radius:50%;background:rgba(201,162,39,.12);border:2px solid rgba(201,162,39,.35)"></div>
      {top(True, "Siguiente paso")}
      <div style="position:absolute;left:64px;top:170px;width:820px">
        <h1 style="font-size:58px">Un proyecto no termina al entregar.<br><span style="color:var(--gold)">Termina al cerrar bien.</span></h1>
        <p class="sub" style="margin:26px 0 30px;max-width:700px">Agenda una reunión de 30 minutos: revisamos tu operación, identificamos las tres brechas de mayor riesgo y te mostramos el plan de 30 días.</p>
        <div style="display:flex;gap:16px"><div class="card" style="min-width:210px"><b style="color:var(--gold)">1 · Diagnóstico</b><br>Conversación y revisión de procesos</div>
        <div class="card" style="min-width:210px"><b style="color:var(--gold)">2 · Propuesta</b><br>Alcance, entregables y calendario</div>
        <div class="card" style="min-width:210px"><b style="color:var(--gold)">3 · Arranque</b><br>Plan de implementación de 30 días</div></div>
        <p style="margin-top:34px;font-size:22px;font-weight:700;color:var(--gold)">{contacto("  ·  ")}</p>
      </div>{pie(10)}</section>''')
    cuerpo = "\n".join(s)
    js = """
const slides=[...document.querySelectorAll('.slide')];let i=0;
function fit(){const k=Math.min(innerWidth/1280,(innerHeight-46)/720);const st=document.getElementById('stage');st.style.transform='scale('+k+')';st.style.marginLeft=((innerWidth-1280*k)/2)+'px';}
function show(n){i=Math.max(0,Math.min(slides.length-1,n));slides.forEach((s,k)=>s.style.display=k===i?'block':'none');document.getElementById('cnt').textContent=(i+1)+' / '+slides.length;history.replaceState(null,'','#'+(i+1));}
addEventListener('keydown',e=>{if(['ArrowRight','PageDown',' '].includes(e.key))show(i+1);if(['ArrowLeft','PageUp'].includes(e.key))show(i-1);if(e.key==='Home')show(0);if(e.key==='End')show(slides.length-1);});
let x0=null;addEventListener('touchstart',e=>x0=e.touches[0].clientX);addEventListener('touchend',e=>{if(x0===null)return;const d=e.changedTouches[0].clientX-x0;if(Math.abs(d)>40)show(i+(d<0?1:-1));x0=null;});
addEventListener('resize',fit);fit();show((parseInt(location.hash.slice(1))||1)-1);
document.getElementById('prev').onclick=()=>show(i-1);document.getElementById('next').onclick=()=>show(i+1);
"""
    return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{CFG["lema"]}</title><style>{CSS}
html,body{{height:100%;overflow:hidden}}
#stage{{width:1280px;height:720px;transform-origin:0 0;position:relative;margin-top:8px}}
.slide{{width:1280px;height:720px;position:absolute;left:0;top:0;overflow:hidden;display:none;border-radius:6px}}
#nav{{position:fixed;left:0;right:0;bottom:0;height:38px;display:flex;justify-content:center;align-items:center;gap:18px;color:#c9d3db;font-size:14px}}
#nav button{{background:rgba(255,255,255,.12);color:#fff;border:0;border-radius:6px;padding:6px 16px;font-size:15px;cursor:pointer}}
@media print{{html,body{{overflow:visible;height:auto;background:#fff}}#nav{{display:none}}#stage{{transform:none!important;margin:0!important;height:auto;width:1280px}}
.slide{{display:block!important;position:relative;page-break-after:always;break-after:page;border-radius:0}}@page{{size:1280px 720px;margin:0}}}}
</style></head><body><div id="stage">{cuerpo}</div>
<div id="nav"><button id="prev">◀</button><span id="cnt"></span><button id="next">▶</button></div><script>{js}</script></body></html>'''


def captura(html_path, png_path, w, h):
    """El modo headless descuenta la barra del navegador: se captura más alto y se recorta al tamaño exacto."""
    subprocess.run([CH, "--headless=new", "--no-sandbox", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=1",
                    f"--window-size={w},{h + 150}", f"--screenshot={png_path}", "file://" + html_path], capture_output=True, timeout=120)
    from PIL import Image
    im = Image.open(png_path)
    im.crop((0, 0, w, h)).save(png_path)


def pdf(html_path, pdf_path):
    subprocess.run([CH, "--headless=new", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer", f"--print-to-pdf={pdf_path}", "file://" + html_path],
                   capture_output=True, timeout=180)


# ----------------------------------------------------------------------------- INFOGRAFÍAS
def info_html(titulo, sub, cuerpo, kicker="Operación bajo control"):
    pie_ = contacto("  ·  ")
    return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><style>{CSS}
body{{width:1080px;height:1350px;overflow:hidden}}
.w{{width:1080px;height:1350px;position:relative;padding:64px;background:linear-gradient(160deg,#071d35 0%,#0B2A4A 50%,#134a6c 100%);color:#fff}}
.w:before{{content:"";position:absolute;right:-180px;top:-180px;width:560px;height:560px;border-radius:50%;background:rgba(201,162,39,.14);border:2px solid rgba(201,162,39,.4)}}
.w h1{{font-size:70px;margin:18px 0 14px;position:relative}}
.w .sub{{font-size:26px;position:relative;max-width:860px}}
.blk{{background:rgba(255,255,255,.09);border:1px solid rgba(255,255,255,.16);border-radius:22px;padding:32px 34px;position:relative}}
.blk h3{{font-size:33px;margin-bottom:8px;color:#fff}}.blk p{{font-size:25px;line-height:1.38;opacity:.92}}
.num{{width:62px;height:62px;border-radius:50%;background:var(--gold);color:var(--navy);font-weight:800;font-size:32px;display:flex;align-items:center;justify-content:center;flex:none}}
.row{{display:flex;gap:20px;align-items:flex-start}}
.cta{{position:absolute;left:0;right:0;bottom:0;background:#051526;padding:30px 64px;border-top:4px solid var(--gold);display:flex;justify-content:space-between;align-items:center}}
.cta b{{font-size:25px;color:var(--gold)}}.cta span{{font-size:21px;opacity:.9}}
.sem{{width:22px;height:22px}}
</style></head><body><div class="w"><div style="position:relative;display:flex;justify-content:space-between;align-items:center">{marca()}<span class="kicker">{kicker}</span></div>
<h1>{titulo}</h1><p class="sub">{sub}</p><div style="margin-top:36px;position:relative">{cuerpo}<div class="blk" style="margin-top:6px;text-align:center;font-size:30px;font-weight:800;color:var(--gold);padding:22px">Menos improvisación. Más trazabilidad.</div></div>
<div class="cta"><div><b>{CFG["lema"]}</b><br><span>Agenda una reunión de 30 minutos</span></div><div style="text-align:right"><span>{pie_}</span></div></div></div></body></html>'''


def pasos_info(items):
    return "".join(f'<div class="blk row" style="margin-bottom:20px"><div class="num">{i}</div><div><h3>{t}</h3><p>{d}</p></div></div>' for i, (t, d) in enumerate(items, 1))


INFOGRAFIAS = [
    ("01_Operacion_bajo_control", "OPERACIÓN BAJO CONTROL", "Cuatro pasos para pasar del desorden a una operación que se puede dirigir.",
     pasos_info([("Ordenar", "Procesos, documentos y responsables en un solo lugar."), ("Planear", "Prioridades, calendario, presupuesto y riesgos."),
                 ("Ejecutar", "Formatos estándar y seguimiento semanal con semáforos."), ("Cerrar", "Entregables validados, cifras conciliadas y expediente completo.")]), "Modelo operativo"),
    ("02_Nomina_sin_sorpresas", "NÓMINA SIN SORPRESAS", "Un flujo claro desde la asistencia hasta el recibo archivado.",
     pasos_info([("Asistencia e incidencias", "Faltas, retardos, vacaciones e incapacidades validados."), ("Cálculo y revisión", "Percepciones, deducciones y neto conciliados."),
                 ("Dispersión", "Pago autorizado con layout y respaldo."), ("Recibos y archivo", "CFDI de nómina y evidencia por periodo.")]), "Nómina"),
    ("03_Cumplimiento_fiscal_e_IMSS", "CUMPLIMIENTO FISCAL E IMSS", "Calendario, acuses y expediente: sin multas, sin improvisar.",
     pasos_info([("Documentos al día", "Constancia, opinión de cumplimiento y contratos vigentes."), ("Cálculo y declaración", "Impuestos y cuotas con cifras conciliadas."),
                 ("Pago y acuse", "Línea de captura pagada y acuse guardado."), ("Expediente por mes", "Todo localizable en minutos si llega una revisión.")]), "Cumplimiento"),
    ("04_Facturacion_y_cobranza", "FACTURACIÓN Y COBRANZA", "Lo que se factura y no se cobra, no es ingreso: es riesgo.",
     pasos_info([("Factura correcta", "Datos fiscales validados y CFDI emitido a tiempo."), ("Seguimiento", "Antigüedad de saldos revisada cada semana."),
                 ("Cobro y aplicación", "Pagos aplicados y estados de cuenta enviados."), ("Conciliación", "Bancos, cartera y contabilidad en el mismo número.")]), "Cobranza"),
    ("05_Compras_y_proyectos", "COMPRAS Y PROYECTOS BAJO CONTROL", "De la requisición al cierre: autorizado, comparado y documentado.",
     pasos_info([("Requisición autorizada", "Nada se compra sin necesidad y responsable."), ("Cuadro comparativo", "Tres cotizaciones, criterio claro y decisión registrada."),
                 ("Orden y recepción", "Conciliación de tres vías: orden, entrada y factura."), ("Avance y presupuesto", "Plan vs. real por proyecto, cada semana.")]), "Compras y proyectos"),
    ("06_Reporte_para_direccion", "REPORTE SEMANAL PARA DIRECCIÓN", "Una página para decidir: avances, pendientes, riesgos y decisiones.",
     '<div class="blk" style="margin-bottom:16px;display:flex;gap:46px;justify-content:center;font-size:26px;font-weight:800"><span><span class="sem v" style="display:inline-block;border-radius:50%;background:#3DAE6B;vertical-align:middle;margin-right:10px"></span>En control</span><span><span class="sem am" style="display:inline-block;border-radius:50%;background:#C9A227;vertical-align:middle;margin-right:10px"></span>Atención</span><span><span class="sem r" style="display:inline-block;border-radius:50%;background:#E8604C;vertical-align:middle;margin-right:10px"></span>Acción</span></div>'
     + pasos_info([("Resumen ejecutivo", "Cómo va la operación y qué se necesita de dirección."), ("Avances y pendientes", "Con responsable, fecha y evidencia."),
                   ("Riesgos y decisiones", "Matriz de riesgos y opciones con recomendación.")]), "Reportes"),
    ("07_Cierra_bien_tus_proyectos", "CIERRA BIEN TUS PROYECTOS", "Un proyecto no termina al entregar. Termina al cerrar bien.",
     pasos_info([("Entregables validados", "Acta de entrega-recepción firmada."), ("Cifras conciliadas", "Presupuesto, costos y facturación cuadran."),
                 ("Pendientes cerrados", "Ninguno sin responsable ni fecha."), ("Expediente final", "Evidencia completa y lecciones aprendidas.")]), "Cierre de proyectos"),
    ("08_Plan_de_30_dias", "PLAN DE IMPLEMENTACIÓN DE 30 DÍAS", "De la solicitud al primer reporte para dirección.",
     pasos_info([("Semana 1 · Ordenar", "Arranque, diagnóstico y documentos."), ("Semana 2 · Planear", "Prioridades, responsables y calendario."),
                 ("Semana 3 · Ejecutar", "Operación con controles y primer reporte."), ("Semana 4 · Cerrar", "Indicadores, procedimientos y plan a 90 días.")]), "Implementación"),
]

MENSAJES = [
    ("Solicitar reunión · empresa industrial", "Hola [NOMBRE], te escribo de {m}. Ayudamos a empresas industriales a ordenar nómina, cumplimiento fiscal/IMSS, compras y cobranza con un tablero semanal para dirección. "
     "¿Te parece si agendamos 30 minutos esta semana para revisar dónde hoy hay más riesgo en [EMPRESA]? Orden para operar. Control para crecer."),
    ("Solicitar reunión · constructora", "Hola [NOMBRE], soy [TU NOMBRE] de {m}. Trabajamos con constructoras para controlar compras, avance vs. presupuesto, pendientes y el cierre documental de cada obra. "
     "Un proyecto no termina al entregar: termina al cerrar bien. ¿Podemos platicar 30 minutos sobre [PROYECTO]?"),
    ("Solicitar reunión · empresa de servicios o comercio", "Hola [NOMBRE], te contacto de {m}. Si hoy la administración de [EMPRESA] depende de archivos sueltos y memoria, podemos dejarte un sistema "
     "de formatos y controles listo en 30 días. ¿Tienes 30 minutos esta semana para mostrártelo?"),
    ("Envío de infografía", "Hola [NOMBRE], te comparto este resumen visual de cómo ordenamos la operación en cuatro pasos. Si algo te hace sentido para [EMPRESA], con gusto lo vemos en una llamada corta."),
    ("Seguimiento a las 48 horas", "Hola [NOMBRE], ¿pudiste ver mi mensaje anterior? Solo quiero saber si te interesa revisar el tablero semanal y el plan de 30 días. Si prefieres otro día, me adapto."),
    ("Seguimiento a los 7 días", "Hola [NOMBRE], te escribo de nuevo con una idea concreta: en una llamada de 20 minutos te mostramos tres brechas típicas (nómina, fiscal/IMSS y cobranza) y cómo se cierran en 30 días. ¿Jueves o viernes?"),
    ("Confirmación de reunión", "Hola [NOMBRE], confirmado: [DÍA] a las [HORA]. Para aprovechar el tiempo, si puedes ten a la mano: organigrama, calendario de nómina y un ejemplo de reporte que hoy recibe dirección. ¡Gracias!"),
    ("Después de la reunión", "Hola [NOMBRE], gracias por tu tiempo. Quedamos en: 1) [ACUERDO 1], 2) [ACUERDO 2]. Te envío hoy la propuesta con alcance, entregables y calendario de 30 días."),
    ("Envío de propuesta", "Hola [NOMBRE], te comparto la propuesta: método de cuatro pasos, entregables, plan de 30 días e inversión. ¿Te parece si la revisamos juntos el [DÍA] para resolver dudas?"),
    ("Cierre / recordatorio", "Hola [NOMBRE], ¿cómo vas con la propuesta? Si quieres, arrancamos el [FECHA] con la semana de diagnóstico para tener el primer reporte para dirección antes de fin de mes."),
    ("Reactivación de prospecto", "Hola [NOMBRE], hace tiempo platicamos sobre ordenar la administración de [EMPRESA]. Hoy contamos con un libro de control semanal con semáforos y reporte para directivos. ¿Retomamos la conversación?"),
]

LEYENDAS = ["Operación bajo control: ordenar, planear, ejecutar y cerrar. ¿Cuál de los cuatro pasos te falta hoy? Escríbenos.",
            "Nómina sin sorpresas: de la asistencia al recibo archivado, con evidencia en cada paso.",
            "Fiscal e IMSS sin improvisar: calendario, acuses y expediente por mes.",
            "Lo que se factura y no se cobra es riesgo. Revisa tu antigüedad de saldos cada semana.",
            "Compras y proyectos bajo control: requisición, comparativo, orden, recepción y cierre.",
            "Una página para decidir: avances, pendientes, riesgos y decisiones requeridas.",
            "Un proyecto no termina al entregar. Termina al cerrar bien.",
            "En 30 días: diagnóstico, plan, operación con controles y primer reporte para dirección."]

GUION = [
    ("Apertura", "Gracias por recibirnos. En 30 minutos queremos mostrarte cómo ordenar la administración de tu empresa para que dirección vea, decida y cierre con evidencia. Al final te pediremos un siguiente paso concreto."),
    ("1 · Portada", "Orden para operar. Control para crecer. Trabajamos con constructoras, industriales y empresas de servicios que ya crecieron más rápido que su administración. Pregunta: ¿cómo te enteras hoy de que un proyecto va mal?"),
    ("2 · Problema e impacto", "Cuando la administración improvisa, la operación paga: documentos dispersos, retrabajos, riesgo fiscal y laboral, decisiones a ciegas. El impacto es sobrecosto, multas, cobranza lenta y proyectos que se entregan pero no se cierran."),
    ("3 · Solución", "Proponemos un sistema listo para operar: 38 formatos, controles con semáforos y un método de cuatro pasos. No vendemos archivos: vendemos orden, trazabilidad y decisiones con información."),
    ("4 · Método", "Ordenar, planear, ejecutar y cerrar. Cada paso tiene responsables, fechas y entregables. Un proyecto no termina al entregar: termina al cerrar bien."),
    ("5 · Procesos", "Mostramos cuatro procesos críticos: nómina, fiscal e IMSS, facturación y cobranza, y compras y proyectos. En todos: responsable, fecha y evidencia. Pregunta: ¿cuál de estos hoy te quita más tiempo?"),
    ("6 · Tablero semanal", "Verde, ámbar y coral. En un minuto dirección sabe qué está en control, qué requiere atención y qué necesita una decisión. El mismo libro alimenta el reporte semanal."),
    ("7 · Qué incluye", "Siete colecciones: contabilidad y finanzas, recursos humanos, inventarios y compras, documentos y actas, contratos, políticas y control directivo. Todo editable con tu logotipo."),
    ("8 · Tabla ejecutiva", "Esta tabla responde tres preguntas: qué se controla, quién responde y qué evidencia queda. Es lo que permite pasar una revisión sin sobresaltos."),
    ("9 · Plan de 30 días", "Semana 1 ordenamos, semana 2 planeamos, semana 3 ejecutamos y semana 4 cerramos con un expediente final y un plan a 90 días. El primer reporte para dirección llega en la semana 3."),
    ("10 · Cierre y próximo paso", "Proponemos una reunión de diagnóstico de 30 minutos con tu equipo administrativo para identificar tres brechas prioritarias y definir fecha de arranque. ¿Qué día de esta semana te acomoda?"),
]


def main():
    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(os.path.join(OUT, "infografias"))
    # presentación
    html = os.path.join(OUT, "Presentacion_Comercial.html")
    open(html, "w", encoding="utf-8").write(deck_html())
    pdf(html, os.path.join(OUT, "Presentacion_Comercial.pdf"))
    # infografías
    tmp = os.path.join(OUT, "_tmp")
    os.makedirs(tmp)
    for nombre, tit, sub, cuerpo, k in INFOGRAFIAS:
        ruta = os.path.join(tmp, nombre + ".html")
        open(ruta, "w", encoding="utf-8").write(info_html(tit, sub, cuerpo, k))
        captura(ruta, os.path.join(OUT, "infografias", nombre + ".png"), 1080, 1350)
    shutil.rmtree(tmp)
    # textos
    m = CFG["marca"].title() if CFG.get("marca") else "[TU MARCA]"
    md = ["# Mensajes de WhatsApp y textos para redes\n", "Sustituye los datos entre [CORCHETES] antes de enviar.\n"]
    for t, x in MENSAJES:
        md.append(f"## {t}\n\n{x.replace('{m}', m)}\n")
    md.append("## Textos para acompañar cada infografía\n")
    for (nombre, tit, *_), l in zip(INFOGRAFIAS, LEYENDAS):
        md.append(f"**{tit.title()}** — {l}\n")
    open(os.path.join(OUT, "Mensajes_WhatsApp.md"), "w", encoding="utf-8").write("\n".join(md))
    md = ["# Guion de presentación (10 láminas)\n", "Estructura: problema, impacto, solución, método, entregables y próximo paso.\n"]
    for t, x in GUION:
        md.append(f"## {t}\n\n{x}\n")
    open(os.path.join(OUT, "Guion_Presentacion.md"), "w", encoding="utf-8").write("\n".join(md))
    from docx import Document
    from docx.shared import Pt, RGBColor
    for nombre, titulo_, bloques in (("Mensajes_WhatsApp", "Mensajes de WhatsApp y textos para redes", [(t, x.replace("{m}", m)) for t, x in MENSAJES] +
                                      [("Texto para " + n_[1].title(), l) for n_, l in zip(INFOGRAFIAS, LEYENDAS)]),
                                     ("Guion_Presentacion", "Guion de presentación comercial (10 láminas)", GUION)):
        d = Document()
        d.styles["Normal"].font.name = "Abadi"
        d.styles["Normal"].font.size = Pt(11)
        h = d.add_heading(titulo_, 0)
        for r in h.runs:
            r.font.color.rgb = RGBColor(0x0B, 0x2A, 0x4A)
        d.add_paragraph("Sustituye los datos entre [CORCHETES] antes de usar." if "Mensajes" in nombre else
                        "Estructura: problema, impacto, solución, método, entregables y próximo paso.")
        for t, x in bloques:
            hh = d.add_heading(t, 2)
            for r in hh.runs:
                r.font.color.rgb = RGBColor(0x0E, 0x7C, 0x8B)
            d.add_paragraph(x)
        d.save(os.path.join(OUT, nombre + ".docx"))
    print("Marketing listo en marketing/")


if __name__ == "__main__":
    main()
