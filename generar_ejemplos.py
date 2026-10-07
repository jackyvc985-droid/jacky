#!/usr/bin/env python3
"""Crea versiones con datos de ejemplo (./ejemplos/) a partir de las plantillas en blanco.
Sirven como vista previa de venta y como demostración para el comprador."""
import datetime as dt
import os
import shutil
from openpyxl import load_workbook
from generar_plantillas import CAJA_F1, INPUT, OP

RAIZ = os.path.dirname(os.path.abspath(__file__))
PLANT = os.path.join(RAIZ, "plantillas")
EJ = os.path.join(RAIZ, "ejemplos")
D = dt.date


def mayus(wb):
    """Nombres, conceptos y textos de captura en MAYÚSCULAS (formato recomendado)."""
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                f = c.fill
                if (isinstance(c.value, str) and not c.value.startswith("=") and f and f.fill_type == "solid"
                        and str(f.fgColor.rgb).endswith(INPUT) and c.value not in ("☐", "☑")):
                    c.value = c.value.upper()


def poner(ws, datos):
    for k, v in datos.items():
        ws[k] = v


def reg(ws, r0, cols, tuplas):
    """Escribe cada tupla en las columnas indicadas (lista de números de columna)."""
    for i, t in enumerate(tuplas):
        for c, v in zip(cols, t):
            if v is not None:
                ws.cell(r0 + i, c, v)


def filas(ws, r0, col0, tuplas):
    for i, t in enumerate(tuplas):
        for j, v in enumerate(t):
            if v is not None:
                ws.cell(r0 + i, col0 + j, v)


def ingresos_gastos(ws):
    poner(ws, {"B4": "Comercializadora Ejemplo S.A. de C.V.", "F4": "Septiembre 2026", "B5": 15000})
    filas(ws, 8, 2, [
        (D(2026, 9, 1), "Venta mostrador", "Ventas", "Efectivo", 18500, None),
        (D(2026, 9, 2), "Pago de renta de oficina", "Renta", "Transferencia", None, 8500),
        (D(2026, 9, 3), "Servicio de consultoría", "Servicios", "Transferencia", 12000, None),
        (D(2026, 9, 5), "Nómina primera quincena", "Nómina", "Transferencia", None, 14200),
        (D(2026, 9, 7), "Papelería e insumos", "Insumos", "Tarjeta", None, 2350.50),
        (D(2026, 9, 9), "Venta mayoreo", "Ventas", "Transferencia", 26400, None),
        (D(2026, 9, 10), "Luz y agua", "Servicios (luz/agua/tel)", "Efectivo", None, 1890),
        (D(2026, 9, 12), "Campaña en redes sociales", "Publicidad", "Tarjeta", None, 3200),
        (D(2026, 9, 14), "Venta mostrador", "Ventas", "Efectivo", 15750, None),
        (D(2026, 9, 15), "Pago de IVA", "Impuestos", "Transferencia", None, 6120),
        (D(2026, 9, 17), "Servicio de mantenimiento", "Servicios", "Transferencia", 9800, None),
        (D(2026, 9, 19), "Combustible de reparto", "Transporte", "Tarjeta", None, 1450),
        (D(2026, 9, 20), "Nómina segunda quincena", "Nómina", "Transferencia", None, 14200),
        (D(2026, 9, 23), "Intereses ganados", "Otros ingresos", "Transferencia", 320, None),
        (D(2026, 9, 25), "Venta mostrador", "Ventas", "Efectivo", 21300, None),
        (D(2026, 9, 28), "Reparaciones menores", "Otros gastos", "Efectivo", None, 1180)])
    extra = [(2, "Venta mostrador", "Ventas", "Efectivo", 9800, None), (4, "Pago de internet y teléfono", "Servicios (luz/agua/tel)", "Tarjeta", None, 1250),
             (6, "Servicio de consultoría", "Servicios", "Transferencia", 7500, None), (8, "Compra de insumos de limpieza", "Insumos", "Efectivo", None, 860),
             (11, "Venta mayoreo", "Ventas", "Transferencia", 31200, None), (13, "Mantenimiento de equipo de cómputo", "Mantenimiento", "Tarjeta", None, 2400),
             (16, "Venta mostrador", "Ventas", "Efectivo", 12400, None), (18, "Combustible", "Transporte", "Tarjeta", None, 1320),
             (21, "Servicio de capacitación", "Servicios", "Transferencia", 14500, None), (22, "Pago ISR provisional", "Impuestos", "Transferencia", None, 4380),
             (24, "Venta mostrador", "Ventas", "Efectivo", 16800, None), (26, "Publicidad impresa", "Publicidad", "Tarjeta", None, 1900),
             (27, "Venta mayoreo", "Ventas", "Transferencia", 22800, None), (29, "Pago de renta bodega", "Renta", "Transferencia", None, 6200)]
    filas(ws, 24, 2, [(D(2026, 9, d_), *resto) for d_, *resto in extra])
    hist = []
    for m_ in range(1, 9):
        hist += [(D(2026, m_, 5), "Venta mostrador", "Ventas", "Efectivo", 14000 + m_ * 1700, None),
                 (D(2026, m_, 12), "Servicio de consultoría", "Servicios", "Transferencia", 8000 + m_ * 900, None),
                 (D(2026, m_, 15), "Nómina quincenal", "Nómina", "Transferencia", None, 13500 + m_ * 100),
                 (D(2026, m_, 20), "Pago de renta", "Renta", "Transferencia", None, 8500),
                 (D(2026, m_, 24), "Insumos y papelería", "Insumos", "Tarjeta", None, 1800 + m_ * 120),
                 (D(2026, m_, 27), "Publicidad en redes", "Publicidad", "Tarjeta", None, 2500)]
    filas(ws, 38, 2, hist)


def caja_chica(ws):
    poner(ws, {"B4": "María González", "F4": D(2026, 9, 1), "B5": 2000})
    filas(ws, 8, 2, [
        (D(2026, 9, 2), "V-001", "Papelería y tóner", "Office Depot", 350),
        (D(2026, 9, 4), "V-002", "Agua purificada", "Garrafones del Norte", 150),
        (D(2026, 9, 8), "V-003", "Mensajería urgente", "Estafeta", 210),
        (D(2026, 9, 11), "V-004", "Material de limpieza", "Súper Mercado", 295),
        (D(2026, 9, 16), "V-005", "Café y consumibles", "Costco", 420)])
    total = sum(x * 1.16 for x in (350, 150, 210, 295, 420))
    ws[f"B{CAJA_F1 + 4}"] = round(2000 - total, 2)


def flujo(ws):
    poner(ws, {"B4": "Comercializadora Ejemplo S.A. de C.V.", "G4": 2026, "B5": 50000})
    ing = [(95, 110, 120, 105, 130, 140, 125, 135, 150, 145, 160, 180), (60, 65, 70, 72, 75, 80, 78, 82, 90, 88, 95, 100),
           (20, 18, 25, 22, 28, 30, 26, 32, 35, 33, 38, 45), (3, 2, 4, 3, 5, 4, 3, 6, 5, 4, 6, 8)]
    egr = [(70, 70, 72, 72, 74, 74, 76, 76, 78, 78, 80, 95), (22, 22, 22, 22, 22, 22, 24, 24, 24, 24, 24, 24),
           (80, 85, 92, 88, 98, 105, 96, 104, 112, 108, 118, 130), (9, 9, 9, 10, 10, 10, 10, 10, 11, 11, 11, 11),
           (18, 20, 24, 22, 26, 28, 25, 28, 32, 30, 34, 40), (6, 6, 8, 6, 8, 8, 6, 8, 10, 8, 10, 14), (5, 4, 6, 5, 6, 7, 5, 7, 8, 6, 8, 9)]
    for i, fila in enumerate(ing):
        for j, v in enumerate(fila):
            ws.cell(9 + i, 2 + j, round(v * 1150))
    for i, fila in enumerate(egr):
        for j, v in enumerate(fila):
            ws.cell(16 + i, 2 + j, v * 1000)


def cxc(ws):
    poner(ws, {"B4": "Comercializadora Ejemplo S.A. de C.V.", "G4": D(2026, 10, 7)})
    reg(ws, 7, [2, 3, 4, 5, 7, 8], [
        ("Grupo Industrial del Norte", "F-1021", D(2026, 6, 15), 30, 48500, 0),
        ("Constructora Altamira", "F-1034", D(2026, 7, 20), 30, 125000, 60000),
        ("Distribuidora La Estrella", "F-1040", D(2026, 8, 12), 15, 32800, 32800),
        ("Servicios Técnicos Omega", "F-1047", D(2026, 8, 28), 30, 76400, 20000),
        ("Hotel Plaza Mayor", "F-1052", D(2026, 9, 10), 30, 54300, 0),
        ("Farmacias San Rafael", "F-1058", D(2026, 9, 22), 15, 18900, 0),
        ("Transportes Veloz", "F-1063", D(2026, 9, 30), 30, 92000, 0),
        ("Alimentos del Valle", "F-1066", D(2026, 10, 3), 45, 41250, 0)])
    reg(ws, 15, [2, 3, 4, 5, 7, 8], [
        ("Colegio Anáhuac del Sur", "F-1069", D(2026, 5, 18), 30, 28700, 0),
        ("Taller Mecánico Rivera", "F-1070", D(2026, 8, 5), 15, 15400, 15400),
        ("Inmobiliaria Torres Díaz", "F-1074", D(2026, 9, 2), 30, 67800, 30000),
        ("Clínica Santa Lucía", "F-1075", D(2026, 9, 18), 30, 83600, 0)])


def inventario(ws):
    poner(ws, {"B4": "Almacén central", "G4": D(2026, 10, 7)})
    reg(ws, 7, [2, 3, 4, 5, 6, 7, 8, 10, 12], [
        ("P-001", "Papel bond carta (caja 5000)", "Papelería", "Caja", 40, 60, 72, 780, 15),
        ("P-002", "Tóner negro HP 85A", "Consumibles", "Pieza", 25, 10, 31, 920, 8),
        ("P-003", "Folder manila carta (100)", "Papelería", "Paquete", 60, 0, 22, 145, 20),
        ("P-004", "Pluma azul (caja 12)", "Papelería", "Caja", 30, 20, 38, 96, 10),
        ("P-005", "Cuaderno profesional", "Papelería", "Pieza", 120, 80, 150, 28, 40),
        ("P-006", "Engrapadora metálica", "Equipo", "Pieza", 12, 0, 4, 185, 6),
        ("P-007", "Cinta adhesiva transparente", "Consumibles", "Rollo", 90, 40, 110, 18, 30),
        ("P-008", "Marcador permanente", "Papelería", "Pieza", 50, 25, 70, 22, 25),
        ("P-009", "Tinta Epson 664 negro", "Consumibles", "Pieza", 18, 0, 18, 210, 6),
        ("P-010", "Silla ergonómica", "Mobiliario", "Pieza", 14, 6, 9, 2450, 4),
        ("P-011", "Escritorio ejecutivo", "Mobiliario", "Pieza", 8, 2, 7, 4890, 3),
        ("P-012", "Clips estándar (caja)", "Papelería", "Caja", 200, 100, 140, 14, 50)])
    reg(ws, 19, [2, 3, 4, 5, 6, 7, 8, 10, 12], [
        ("P-013", "Calculadora científica", "Equipo", "Pieza", 20, 10, 18, 320, 6),
        ("P-014", "Archivero metálico 4 gavetas", "Mobiliario", "Pieza", 6, 0, 5, 3290, 2),
        ("P-015", "Etiquetas adhesivas (caja)", "Papelería", "Caja", 60, 30, 70, 54, 20),
        ("P-016", "Sobre manila carta (100)", "Papelería", "Paquete", 45, 0, 38, 120, 12),
        ("P-017", "Lámpara de escritorio LED", "Equipo", "Pieza", 15, 5, 12, 410, 5),
        ("P-018", "Pizarrón blanco 90x120", "Mobiliario", "Pieza", 5, 0, 5, 1480, 2)])


def presupuesto(ws):
    poner(ws, {"B4": "Comercializadora Ejemplo S.A. de C.V.", "G4": 2026})
    base = [(220, 1.0), (120, 1.0), (62, 1.0), (22, 1.0), (9, 1.0), (14, 1.0), (6, 1.0), (4, 1.0), (16, 1.0), (7, 1.0)]
    var = [1.06, 0.97, 1.04, 1.12, 1.02, 0.93, 1.15, 1.08, 1.01, 1.2]
    for i, (b, _) in enumerate(base):
        for m in range(12):
            ws.cell(8 + 2 * i, 2 + m, b * 1000)
            if m < 9:
                ws.cell(9 + 2 * i, 2 + m, round(b * 1000 * (var[i] + (m % 3 - 1) * 0.02)))


def nomina(ws):
    poner(ws, {"B4": "Comercializadora Ejemplo S.A. de C.V.", "G4": "1ª quincena de octubre 2026"})
    reg(ws, 8, [2, 3, 4, 5, 8, 10], [
        ("Ana López Ramírez", 520, 15, 0, 0, 1420),
        ("Carlos Pérez Soto", 410, 15, 4, 300, 910),
        ("Laura Méndez Cruz", 680, 15, 0, 500, 1880),
        ("José Hernández Ruiz", 350, 14, 6, 0, 640),
        ("Patricia Núñez Vega", 460, 15, 2, 250, 1030),
        ("Miguel Ángel Torres", 395, 15, 0, 0, 790)])
    reg(ws, 14, [2, 3, 4, 5, 8, 10], [
        ("Sofía Ramos Díaz", 430, 15, 0, 0, 840), ("Luis Castillo Mora", 365, 15, 3, 150, 720),
        ("Daniela Fuentes Ibarra", 590, 15, 0, 400, 1510), ("Ricardo Salinas Paz", 480, 14, 2, 0, 960)])


def cotizacion(ws):
    poner(ws, {"B3": "Tu Empresa S.A. de C.V.", "F3": "COT-2026-0142", "B4": "TEM260101AB1", "F4": D(2026, 10, 7),
               "B5": "Constructora Altamira S.A.", "B6": "Ing. Roberto Salas", "F6": "55 1234 5678"})
    filas(ws, 9, 2, [
        ("SRV-01", "Diagnóstico administrativo", "Servicio", 1, 18500),
        ("SRV-02", "Implementación de control de inventarios", "Servicio", 1, 32000),
        ("CAP-01", "Capacitación al personal (8 horas)", "Curso", 2, 6500),
        ("SOP-01", "Soporte mensual (3 meses)", "Mes", 3, 4200)])
    ws["G25"] = 1500


def orden_compra(ws):
    poner(ws, {"B3": "Tu Empresa S.A. de C.V.", "F3": "OC-0087", "B4": "TEM260101AB1", "F4": D(2026, 10, 7),
               "B5": "Papelera del Centro S.A.", "F5": D(2026, 10, 14), "B6": "Lic. Sandra Ríos",
               "F6": "Crédito 30 días", "B7": "Av. Reforma 245, Col. Centro, CDMX"})
    filas(ws, 10, 2, [
        ("P-001", "Papel bond carta (caja 5000)", "Caja", 20, 780),
        ("P-002", "Tóner negro HP 85A", "Pieza", 10, 920),
        ("P-004", "Pluma azul (caja 12)", "Caja", 15, 96),
        ("P-007", "Cinta adhesiva transparente", "Rollo", 60, 18)])


def evaluacion(ws):
    poner(ws, {"B4": "Ana López Ramírez", "E4": "Coordinadora administrativa", "B5": "Administración",
               "E5": "Enero - Junio 2026", "B6": "Lic. Jorge Medina", "E6": D(2026, 7, 3)})
    for r, v in zip(range(9, 17), (5, 4, 4, 5, 4, 4, 5, 5)):
        ws.cell(r, 3, v)
    ws["B21"] = "Excelente organización, cumple fechas límite y apoya a su equipo."
    ws["B26"] = "Fortalecer delegación de tareas; curso de liderazgo en el segundo semestre."
    ws["F9"] = "Cero errores en cierres mensuales."
    ws["F10"] = "Cumplió 96% de objetivos."


def asistencia(ws):
    poner(ws, {"C4": "Comercializadora Ejemplo S.A. de C.V.", "C5": "Septiembre 2026"})
    nombres = ["Ana López Ramírez", "Carlos Pérez Soto", "Laura Méndez Cruz", "José Hernández Ruiz",
               "Patricia Núñez Vega", "Miguel Ángel Torres", "Sofía Ramos Díaz", "Luis Castillo Mora"]
    for i, n in enumerate(nombres):
        ws.cell(7 + i, 2, n)
        for dia in range(1, 31):
            fecha = D(2026, 9, dia)
            cod = "D" if fecha.weekday() >= 5 else "A"
            if (i * 7 + dia) % 23 == 0 and cod == "A":
                cod = "R"
            if (i * 5 + dia) % 29 == 0 and cod == "A":
                cod = "F"
            if i == 2 and 14 <= dia <= 18 and cod == "A":
                cod = "V"
            ws.cell(7 + i, 2 + dia, cod)
    mas = ["Sofía Ramos Díaz", "Luis Castillo Mora", "Daniela Fuentes Ibarra", "Ricardo Salinas Paz", "Fernanda Ochoa León", "Jorge Medina Rangel"]
    for k, n in enumerate(mas):
        i = 8 + k
        ws.cell(7 + i, 2, n)
        for dia in range(1, 31):
            fecha = D(2026, 9, dia)
            cod = "D" if fecha.weekday() >= 5 else "A"
            if (i * 7 + dia) % 19 == 0 and cod == "A":
                cod = "R"
            if (i * 5 + dia) % 31 == 0 and cod == "A":
                cod = "F"
            if i == 11 and 7 <= dia <= 11 and cod == "A":
                cod = "I"
            ws.cell(7 + i, 2 + dia, cod)


def vacaciones(ws):
    poner(ws, {"B4": "Comercializadora Ejemplo S.A. de C.V.", "G4": D(2026, 10, 7)})
    reg(ws, 7, [2, 3, 6, 9], [
        ("Ana López Ramírez", D(2019, 3, 4), 6, 520), ("Carlos Pérez Soto", D(2023, 8, 14), 0, 410),
        ("Laura Méndez Cruz", D(2016, 1, 11), 10, 680), ("José Hernández Ruiz", D(2025, 2, 3), 0, 350),
        ("Patricia Núñez Vega", D(2021, 6, 21), 8, 460)])

def cuadro(ws):
    poner(ws, {"B4": "REQ-2026-031", "F4": D(2026, 10, 5), "E6": "Papelera del Centro", "G6": "Office Max",
               "I6": "Distribuidora Pluma", "K6": "Comercial Rex"})
    datos = [("Papel bond carta (caja)", "Caja", 20, 780, 805, 790, None), ("Tóner HP 85A", "Pieza", 10, 920, 890, 950, 905),
             ("Pluma azul (caja 12)", "Caja", 15, 96, 102, 90, 99), ("Cinta adhesiva", "Rollo", 60, 18, 17.5, 19, 18)]
    for i, (d_, u, c, a, b, e, f) in enumerate(datos):
        r = 8 + i
        ws.cell(r, 2, d_), ws.cell(r, 3, u), ws.cell(r, 4, c)
        for col, v in ((5, a), (7, b), (9, e), (11, f)):
            if v is not None:
                ws.cell(r, col, v)
    ws["F26"] = "Crédito 30 días"; ws["H26"] = "Contado"; ws["J26"] = "Crédito 15 días"; ws["L26"] = "Crédito 30 días"
    ws["F27"] = 3; ws["H27"] = 5; ws["J27"] = 2; ws["L27"] = 7


def activos(ws):
    poner(ws, {"B4": "Comercializadora Ejemplo S.A. de C.V.", "G4": D(2026, 10, 7)})
    reg(ws, 7, [2, 3, 4, 5, 6, 12], [
        ("AF-001", "Laptop Dell Latitude 5440", "Equipo de cómputo", D(2024, 2, 15), 28900, "Administración"),
        ("AF-002", "Escritorio ejecutivo", "Mobiliario y equipo de oficina", D(2022, 6, 1), 4890, "Dirección"),
        ("AF-003", "Camioneta Nissan NP300", "Automóviles", D(2023, 1, 10), 389000, "Reparto"),
        ("AF-004", "Impresora multifuncional", "Equipo de cómputo", D(2025, 3, 20), 12800, "Recepción"),
        ("AF-005", "Aire acondicionado 2 ton", "Maquinaria y equipo", D(2021, 5, 12), 21500, "Sala de juntas")])

def directorio(ws):
    poner(ws, {"B4": "Comercializadora Ejemplo S.A. de C.V."})
    filas(ws, 7, 2, [
        ("Cliente", "Constructora Altamira S.A. de C.V.", "CAL150312AB4", "Ing. Roberto Salas", "55 1234 5678", "rsalas@altamira.mx", 30, 150000, "Av. Insurgentes 1020, CDMX"),
        ("Cliente", "Hotel Plaza Mayor S.A.", "HPM110920XY2", "Lic. Elena Cruz", "33 4455 6677", "compras@plazamayor.mx", 30, 80000, "Guadalajara, Jal."),
        ("Proveedor", "Papelera del Centro S.A.", "PCE091105KL7", "Lic. Sandra Ríos", "55 9988 7766", "ventas@papelcentro.mx", 30, None, "Centro, CDMX"),
        ("Proveedor", "Distribuidora Pluma S. de R.L.", "DPL130228MN3", "Sr. Alberto Mora", "81 2233 4455", "amora@pluma.mx", 15, None, "Monterrey, N.L."),
        ("Ambos", "Servicios Técnicos Omega", "STO120714QR9", "Ing. Daniel Vega", "442 118 2020", "contacto@omega.mx", 45, 60000, "Querétaro, Qro.")])


def conciliacion(ws):
    poner(ws, {"B4": "BBVA", "B5": "0123456789", "B6": "Septiembre 2026", "B7": D(2026, 9, 30),
               "C10": 184250.40, "C13": 0, "C17": 188280.15, "C20": 0, "C21": 0})
    pt = ws.parent["PARTIDAS"]
    reg(pt, 5, [2, 3, 4], [(D(2026, 9, 30), "Depósito de ventas 30-sep", 12500)])
    reg(pt, 5, [8, 9, 10], [(D(2026, 9, 28), "Cheque 1045 Papelera del Centro", 5400),
                            (D(2026, 9, 29), "Cheque 1046 Distribuidora Pluma", 2900.25)])
    reg(pt, 5, [14, 15, 16, 17], [(D(2026, 9, 30), "Intereses ganados", "ABONO", 320),
                                  (D(2026, 9, 30), "Comisión por manejo de cuenta", "CARGO", 150)])


def recibo(ws):
    poner(ws, {"B3": "R-0215", "E3": D(2026, 10, 7), "B4": 15000, "B5": "Constructora Altamira S.A. de C.V.",
               "B6": "Quince mil pesos 00/100 M.N.", "B7": "Pago parcial de la factura F-1034 (servicios de mantenimiento)",
               "B8": "Transferencia", "E8": "SPEI 884520", "B10": 60000, "B11": 15000})


def kardex(ws):
    poner(ws, {"B4": "Papel bond carta (caja 5000)", "G4": "P-001", "B5": "Caja"})
    reg(ws, 9, [2, 3, 4, 5, 6, 8], [
        (D(2026, 9, 1), "INV-INICIAL", "Inventario inicial", 40, 760, None),
        (D(2026, 9, 3), "FAC-5521", "Compra a Papelera del Centro", 60, 780, None),
        (D(2026, 9, 5), "REM-0311", "Venta mostrador", None, None, 18),
        (D(2026, 9, 12), "REM-0328", "Venta mayoreo", None, None, 30),
        (D(2026, 9, 18), "FAC-5560", "Compra a Distribuidora Pluma", 40, 795, None),
        (D(2026, 9, 24), "REM-0351", "Venta mostrador", None, None, 24)])


def control_operativo(wb):
    wb["TABLERO"]["H8"] = D(2026, 10, 7)
    a = wb["ACTIVIDADES"]
    lun = D(2026, 10, 5)
    reg(a, OP["ACT"][0], [2, 3, 4, 5, 6, 7, 8, 9, 10, 13], [
        (lun, "Conciliar facturas de proveedores de septiembre", "Administración", "Finanzas", "Alta", D(2026, 10, 6), D(2026, 10, 6), "Terminado", 1, "Reporte de conciliación"),
        (lun, "Integrar expediente de compras de la Nave Toluca", "Nave Toluca fase 2", "Compras", "Alta", D(2026, 10, 7), None, "En proceso", 0.7, "Expediente digital"),
        (lun, "Presentar propuesta de cierre a Constructora Altamira", "Remodelación oficinas", "Proyectos", "Alta", D(2026, 10, 9), None, "En proceso", 0.5, "Presentación"),
        (lun, "Calcular nómina de la segunda quincena de septiembre", "Nómina", "Recursos humanos", "Alta", D(2026, 10, 1), D(2026, 10, 1), "Terminado", 1, "Layout de dispersión"),
        (lun, "Presentar declaración mensual de IVA", "Cumplimiento fiscal", "Finanzas", "Alta", D(2026, 10, 5), D(2026, 10, 5), "Terminado", 1, "Acuse SAT"),
        (lun, "Enviar facturas pendientes de timbrar", "Facturación", "Administración", "Media", D(2026, 10, 2), None, "Pendiente", 0, ""),
        (lun, "Dar seguimiento a cobranza mayor a 60 días", "Cobranza", "Finanzas", "Alta", D(2026, 10, 6), None, "En proceso", 0.4, "Bitácora de llamadas"),
        (lun, "Validar altas y bajas de IMSS del mes", "Cumplimiento IMSS", "Recursos humanos", "Media", D(2026, 10, 8), None, "En proceso", 0.6, "Cédula IDSE"),
        (lun, "Actualizar programa de obra Planta Querétaro", "Cierre de obra Querétaro", "Proyectos", "Alta", D(2026, 10, 7), None, "Bloqueado", 0.3, "Programa actualizado"),
        (lun, "Solicitar cotizaciones de acero estructural", "Nave Toluca fase 2", "Compras", "Media", D(2026, 10, 9), None, "Pendiente", 0, ""),
        (lun, "Elaborar reporte semanal para dirección", "Reportes", "Administración", "Alta", D(2026, 10, 9), None, "Pendiente", 0.2, "Reporte semanal"),
        (lun, "Documentar el proceso de alta de proveedores", "ERP administrativo", "Administración", "Baja", D(2026, 10, 14), None, "En proceso", 0.35, "Procedimiento"),
        (lun, "Conciliar caja chica y viáticos", "Control financiero", "Finanzas", "Media", D(2026, 10, 5), D(2026, 10, 6), "Terminado", 1, "Arqueo firmado"),
        (lun, "Capacitar a usuarios en el nuevo módulo de compras", "ERP administrativo", "Operaciones", "Media", D(2026, 10, 12), None, "Pendiente", 0, ""),
        (lun, "Cerrar expediente fotográfico de la obra", "Cierre de obra Querétaro", "Operaciones", "Alta", D(2026, 10, 3), None, "En proceso", 0.8, "Álbum de evidencias"),
        (lun, "Revisar contratos de subcontratistas", "Nave Toluca fase 2", "Dirección", "Media", D(2026, 10, 8), None, "En proceso", 0.5, "Matriz de contratos"),
        (lun, "Autorizar compras mayores a $50,000", "Compras", "Dirección", "Alta", D(2026, 10, 7), D(2026, 10, 7), "Terminado", 1, "Autorizaciones firmadas"),
        (lun, "Actualizar indicadores semanales", "Tablero", "Administración", "Media", D(2026, 10, 9), None, "Pendiente", 0, "")])
    p = wb["PROYECTOS"]
    reg(p, OP["PRO"][0], [2, 3, 4, 5, 6, 8, 9, 10, 14], [
        ("Nave industrial Toluca - fase 2", "Grupo Industrial del Norte", "Proyectos", D(2026, 7, 1), D(2026, 12, 15), 0.42, 4800000, 2350000, "En proceso"),
        ("Remodelación de oficinas corporativas", "Constructora Altamira", "Operaciones", D(2026, 8, 10), D(2026, 11, 20), 0.55, 1250000, 640000, "En proceso"),
        ("Implementación de ERP administrativo", "Administración", "Administración", D(2026, 9, 1), D(2026, 12, 31), 0.18, 780000, 210000, "En proceso"),
        ("Auditoría de cumplimiento IMSS", "Recursos humanos", "Recursos humanos", D(2026, 9, 15), D(2026, 10, 30), 0.70, 120000, 95000, "En proceso"),
        ("Cierre de obra Planta Querétaro", "Operaciones", "Proyectos", D(2026, 5, 4), D(2026, 10, 9), 0.78, 3100000, 3020000, "En proceso")])
    e = wb["PENDIENTES"]
    reg(e, OP["PEN"][0], [2, 3, 4, 5, 6, 7, 8, 11], [
        ("Entregar acta de entrega-recepción firmada", "Cliente Altamira", "Proyectos", "Alta", D(2026, 9, 28), D(2026, 10, 6), "En proceso", "Falta firma del residente"),
        ("Definir proveedor de acero estructural", "Junta semanal", "Compras", "Alta", D(2026, 9, 30), D(2026, 10, 9), "Pendiente", "Esperando 3a cotización"),
        ("Regularizar pago de cuotas IMSS de agosto", "Auditoría", "Recursos humanos", "Alta", D(2026, 9, 22), D(2026, 10, 2), "En proceso", "Pendiente comprobante bancario"),
        ("Actualizar contrato marco con proveedor de concreto", "Dirección", "Dirección", "Media", D(2026, 9, 15), D(2026, 10, 20), "Pendiente", ""),
        ("Enviar estado de cuenta a clientes con saldo vencido", "Finanzas", "Finanzas", "Media", D(2026, 10, 1), D(2026, 10, 8), "En proceso", ""),
        ("Cargar facturas de gastos de septiembre", "Contabilidad", "Administración", "Baja", D(2026, 10, 2), D(2026, 10, 12), "Pendiente", ""),
        ("Resolver diferencia de inventario en bodega 2", "Operaciones", "Operaciones", "Media", D(2026, 9, 25), D(2026, 10, 5), "Terminado", "Ajuste autorizado")])
    r = wb["RIESGOS"]
    reg(r, OP["RIE"][0], [2, 3, 4, 5, 8, 9, 10, 11], [
        ("Retraso en entrega de acero estructural", "Nave Toluca fase 2", 4, 5, "Compras", "Cotizar con dos proveedores alternos y fijar penalización", D(2026, 10, 15), "En mitigación"),
        ("Sobrecosto en cierre de obra Querétaro", "Cierre de obra Querétaro", 4, 4, "Proyectos", "Congelar alcance y autorizar cambios solo por escrito", D(2026, 10, 12), "Abierto"),
        ("Observaciones del IMSS por diferencias en cuotas", "Cumplimiento IMSS", 3, 4, "Recursos humanos", "Conciliar cédula de determinación contra nómina", D(2026, 10, 30), "En mitigación"),
        ("Cobranza vencida mayor a 90 días", "Cobranza", 3, 3, "Finanzas", "Plan de pagos y suspensión de crédito", D(2026, 10, 20), "Abierto"),
        ("Rotación de personal clave en administración", "ERP administrativo", 2, 4, "Recursos humanos", "Documentar procesos y capacitar suplentes", D(2026, 11, 15), "Abierto"),
        ("Falla en el respaldo de información", "Tecnología", 2, 5, "Administración", "Respaldo diario en nube y prueba mensual de restauración", D(2026, 10, 25), "Abierto"),
        ("Inflación en materiales de construcción", "Nave Toluca fase 2", 3, 3, "Compras", "Cláusula de ajuste de precios en contratos", D(2026, 11, 1), "Abierto"),
        ("Errores en timbrado de facturas", "Facturación", 2, 2, "Administración", "Validación previa con lista de verificación", D(2026, 10, 18), "Abierto")])
    d = wb["DECISIONES"]
    reg(d, OP["DEC"][0], [2, 3, 4, 5, 6, 7, 8], [
        ("Autorizar ampliación de presupuesto de cierre de obra", "Sobrecosto de 4.8% por obra adicional solicitada por el cliente", "A) Autorizar $150,000; B) Negociar costo con cliente; C) Absorber", "Opción B con respaldo de orden de cambio", "Dirección", D(2026, 10, 6), "Pendiente"),
        ("Seleccionar proveedor de acero estructural", "Tres cotizaciones con diferencias de precio y tiempo de entrega", "A) Proveedor local; B) Importado; C) Dividir pedido", "Opción C para mitigar retraso", "Dirección", D(2026, 10, 9), "Pendiente"),
        ("Contratar apoyo temporal para cierre fiscal", "Carga de trabajo del cierre de septiembre", "A) Contratar; B) Horas extra", "Opción A por 4 semanas", "Dirección", D(2026, 10, 3), "Aprobada"),
        ("Cambiar proceso de autorización de compras mayores a $50,000", "Compras sin soporte completo", "A) Doble firma; B) Comité semanal", "Opción A", "Dirección", D(2026, 10, 12), "Pendiente")])
    i = wb["INDICADORES"]
    reg(i, OP["IND"][0], [2, 3, 4, 5, 6, 7, 8, 9, 10, 11], [
        ("% de actividades cumplidas en tiempo", "%", "Mayor es mejor", 90, 78, 82, 85, 84, 88, 86),
        ("Facturación semanal", "MXN", "Mayor es mejor", 300000, 255000, 280000, 310000, 295000, 330000, 318000),
        ("Cobranza semanal", "MXN", "Mayor es mejor", 280000, 190000, 210000, 240000, 255000, 262000, 270000),
        ("Días promedio de cobro", "días", "Menor es mejor", 35, 48, 46, 44, 43, 41, 40),
        ("Compras autorizadas en tiempo", "%", "Mayor es mejor", 95, 88, 90, 92, 91, 94, 96),
        ("Incidencias de nómina", "casos", "Menor es mejor", 3, 7, 6, 5, 5, 4, 3),
        ("Pendientes vencidos", "casos", "Menor es mejor", 5, 12, 11, 9, 8, 7, 6),
        ("Avance promedio de proyectos", "%", "Mayor es mejor", 60, 38, 41, 44, 46, 49, 52)])


EJEMPLOS = {
    "1_Contabilidad_y_Finanzas/01_Control_de_Ingresos_y_Gastos": ingresos_gastos,
    "7_Control_Directivo/01_Control_Operativo_Semanal": control_operativo,
    "1_Contabilidad_y_Finanzas/02_Control_de_Caja_Chica": caja_chica,
    "1_Contabilidad_y_Finanzas/03_Flujo_de_Efectivo_Anual": flujo,
    "1_Contabilidad_y_Finanzas/04_Cuentas_por_Cobrar": cxc,
    "1_Contabilidad_y_Finanzas/05_Conciliacion_Bancaria": conciliacion,
    "1_Contabilidad_y_Finanzas/06_Presupuesto_Anual": presupuesto,
    "1_Contabilidad_y_Finanzas/07_Control_de_Activos_Fijos": activos,
    "1_Contabilidad_y_Finanzas/08_Directorio_Clientes_y_Proveedores": directorio,
    "2_Recursos_Humanos/01_Control_de_Asistencia": asistencia,
    "2_Recursos_Humanos/02_Nomina_Simplificada": nomina,
    "2_Recursos_Humanos/03_Control_de_Vacaciones": vacaciones,
    "2_Recursos_Humanos/04_Evaluacion_de_Desempeno": evaluacion,
    "3_Inventarios_y_Compras/01_Control_de_Inventario": inventario,
    "3_Inventarios_y_Compras/02_Kardex_Costo_Promedio": kardex,
    "3_Inventarios_y_Compras/03_Orden_de_Compra": orden_compra,
    "3_Inventarios_y_Compras/04_Cuadro_Comparativo_de_Cotizaciones": cuadro,
    "4_Documentos_y_Actas/01_Cotizacion": cotizacion,
    "4_Documentos_y_Actas/02_Recibo_de_Pago": recibo,
}


def generar():
    shutil.rmtree(EJ, ignore_errors=True)
    for ruta, fn in EJEMPLOS.items():
        wb = load_workbook(os.path.join(PLANT, ruta + ".xlsx"))
        fn(wb["FORMATO"] if "FORMATO" in wb.sheetnames else wb)
        mayus(wb)
        destino = os.path.join(EJ, ruta + "_EJEMPLO.xlsx")
        os.makedirs(os.path.dirname(destino), exist_ok=True)
        wb.save(destino)
    print(f"Ejemplos: {len(EJEMPLOS)} archivos en ejemplos/")


if __name__ == "__main__":
    generar()
